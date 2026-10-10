"""Importación de CSV/XLSX bancarios: lectura, detección de columnas y parseo."""
from __future__ import annotations

import csv
import hashlib
import io
import re
from collections import Counter
from datetime import date, datetime

from ..schemas import ColumnMapping
from .categorizer import norm

DATE_FORMATS = ["%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y", "%Y/%m/%d", "%d-%m-%y", "%d.%m.%y", "%m/%d/%Y"]
HDR = {
    "date": ["fecha de la operacion", "fecha de operacion", "fecha operacion", "fecha valor", "fecha", "f. operacion", "f.valor", "date", "booking date", "f. valor"],
    "description": ["concepto", "descripcion", "detalle", "movimiento", "description", "details", "observaciones", "beneficiario"],
    "amount": ["importe", "amount", "cantidad", "monto"],
    "debit": ["cargo", "debe", "debit", "gasto", "salida"],
    "credit": ["abono", "haber", "credit", "ingreso", "entrada"],
    "balance": ["saldo", "balance"],  # saldo tras el movimiento (opcional; permite cuadrar la cuenta)
}


def read_table(filename: str, data: bytes) -> list[list[str]]:
    """Devuelve filas de texto. XLSX vía openpyxl; CSV con detección de codificación y delimitador."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else (c.isoformat() if isinstance(c, (datetime, date)) else str(c)) for c in r])
        return [r for r in rows if any(c.strip() for c in r)]
    if filename.lower().endswith(".xls"):
        raise ValueError("Formato .xls antiguo no soportado: guarda el archivo como .xlsx o .csv")
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4000]
    counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
    delim = max(counts, key=counts.get)
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    return [[c.strip() for c in r] for r in rows if any(c.strip() for c in r)]


def parse_number(s: str, decimal: str | None = None) -> float | None:
    s = (s or "").strip().replace("€", "").replace("EUR", "").replace("\xa0", "").replace(" ", "")
    if not s:
        return None
    neg = s.startswith("(") and s.endswith(")") or s.endswith("-")
    s = s.strip("()").rstrip("-")
    if not re.fullmatch(r"[+-]?[\d.,]+", s):
        return None
    if decimal is None:
        decimal = "," if s.rfind(",") > s.rfind(".") else "."
        if "," in s and "." not in s and re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+", s):
            decimal = "."  # 1,234,567 => miles
    if decimal == ",":
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        v = float(s)
    except ValueError:
        return None
    return -abs(v) if neg else v


def parse_date(s: str, fmt: str | None) -> date | None:
    s = (s or "").strip().split(" ")[0].split("T")[0]
    for f in ([fmt] if fmt else DATE_FORMATS):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def _detect_date_format(vals: list[str]) -> str | None:
    vals = [v for v in vals if v.strip()][:30]
    for f in DATE_FORMATS:
        if vals and all(parse_date(v, f) for v in vals):
            return f
    return None


def detect(rows: list[list[str]]) -> dict:
    """Localiza la fila de cabecera y propone el mapeo de columnas."""
    header_idx = 0
    for i, r in enumerate(rows[:15]):
        cells = [norm(c).lower() for c in r]
        if sum(1 for c in cells if c and any(k in c for ks in HDR.values() for k in ks)) >= 2:
            header_idx = i
            break
    header = rows[header_idx]
    body = [r + [""] * (len(header) - len(r)) for r in rows[header_idx + 1:]]
    cols = {h: [r[j] for r in body[:50] if j < len(r)] for j, h in enumerate(header) if h}
    mapping: dict[str, str | None] = {k: None for k in HDR}
    for key, kws in HDR.items():
        for kw in kws:
            hit = next((h for h in header if h and h not in mapping.values() and kw in norm(h).lower()), None)
            if hit:
                mapping[key] = hit
                break
    # por contenido, si la cabecera no basta
    if not mapping["date"]:
        mapping["date"] = next((h for h, v in cols.items() if _detect_date_format(v)), None)
    if not mapping["amount"] and not (mapping["debit"] and mapping["credit"]):
        mapping["amount"] = next((h for h, v in cols.items() if h not in mapping.values() and v and sum(parse_number(x) is not None for x in v) / len(v) > 0.8 and not _detect_date_format(v)), None)
    if not mapping["description"]:
        mapping["description"] = max((h for h in cols if h not in mapping.values()), key=lambda h: sum(len(x) for x in cols[h]), default=None)
    fmt = _detect_date_format(cols.get(mapping["date"] or "", []))
    nums = [x for x in cols.get(mapping["amount"] or "", []) if x.strip()]
    commas = sum(1 for x in nums if "," in x and x.rfind(",") > x.rfind("."))
    decimal = "," if nums and commas >= len(nums) / 2 else "."
    return {"header_row": header_idx, "columns": header, "mapping": mapping, "date_format": fmt, "decimal": decimal, "sample": body[:10], "total_rows": len(body)}


def parse_rows(rows: list[list[str]], header_row: int, m: ColumnMapping) -> tuple[list[dict], list[str]]:
    header = rows[header_row]
    idx = {h: i for i, h in enumerate(header)}
    errors, out = [], []

    def cell(r, name):
        i = idx.get(name) if name else None
        return r[i] if i is not None and i < len(r) else ""

    for n, r in enumerate(rows[header_row + 1:], start=header_row + 2):
        d = parse_date(cell(r, m.date), m.date_format)
        if m.amount:
            amt = parse_number(cell(r, m.amount), m.decimal)
        else:
            deb, cre = parse_number(cell(r, m.debit), m.decimal) or 0.0, parse_number(cell(r, m.credit), m.decimal) or 0.0
            amt = cre - abs(deb) if (deb or cre) else None
        if d is None or amt is None:
            errors.append(f"Fila {n}: fecha o importe no válidos")
            continue
        if m.invert_sign:
            amt = -amt
        if amt == 0:
            continue
        out.append({"date": d, "amount": round(abs(amt), 2), "type": "income" if amt > 0 else "expense", "description": cell(r, m.description)[:500],
                    "balance": parse_number(cell(r, m.balance), m.decimal) if m.balance else None})
    return out, errors


def closing_balance(items: list[dict]) -> tuple[float, date] | None:
    """Saldo tras el último movimiento cronológico (el extracto puede venir en orden ascendente o descendente)."""
    with_bal = [it for it in items if it.get("balance") is not None]
    if not with_bal:
        return None
    last = with_bal[0] if with_bal[0]["date"] >= with_bal[-1]["date"] else with_bal[-1]
    return round(last["balance"], 2), last["date"]


def assign_hashes(items: list[dict], account_id: int) -> None:
    """Hash estable; si hay dos movimientos idénticos el mismo día, se distinguen por ocurrencia."""
    seen: Counter = Counter()
    for it in items:
        base = f"{account_id}|{it['date']}|{it['amount']:.2f}|{it['type']}|{norm(it['description'])}"
        seen[base] += 1
        it["hash"] = hashlib.sha1(f"{base}|{seen[base]}".encode()).hexdigest()
